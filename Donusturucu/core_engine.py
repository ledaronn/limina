# core_engine.py — donusturme mantigi. Arayuzden bagimsiz, MCP sunucusundan da
# dogrudan GUI'den de cagrilir.
from __future__ import annotations

import difflib
import io
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from PIL import Image

try:
    import pymupdf as fitz
    from pptx import Presentation
    from pptx.util import Inches
except ImportError:
    fitz = None
    Presentation = None

# LibreOffice kurulumunun tipik konumlari. PATH'te 'soffice' varsa once o denenir.
SOFFICE_ADAYLAR = [
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
]
LIBREOFFICE_TIMEOUT = 120  # saniye


class DonusturmeHatasi(Exception):
    """Beklenen (kullanicinin duzeltebilecegi) donusum hatalari icin."""


def _soffice_yolu() -> str | None:
    yol = shutil.which("soffice")
    if yol:
        return os.path.realpath(yol)
    for aday in SOFFICE_ADAYLAR:
        if os.path.isfile(aday):
            return os.path.realpath(aday)
    return None


class DocumentConverter:
    """
    PPTX/PPT/DOCX/MD/TXT -> PDF, PDF -> PPTX veya PNG, PNG/JPG -> PDF.
    MD/TXT UTF-8 duz metin olarak aktarilir; Markdown isaretleri korunur.
    Baska hicbir donusum desteklenmez.
    """

    CONVERSION_MAP = {
        ".pptx": [".pdf"],
        ".ppt":  [".pdf"],
        ".docx": [".pdf"],
        ".md":   [".pdf"],
        ".txt":  [".pdf"],
        ".pdf":  [".pptx", ".png"],
        ".png":  [".pdf"],
        ".jpg":  [".pdf"],
        ".jpeg": [".pdf"],
    }

    # Toplu hedef secimi acilir menusu icin.
    ALL_TARGET_FORMATS = ["PDF", "PPTX", "PNG"]

    def get_supported_targets(self, file_path: str) -> list:
        ext = os.path.splitext(file_path)[1].lower()
        return self.CONVERSION_MAP.get(ext, [])

    # --- LibreOffice: PPTX/PPT/DOCX/MD/TXT -> PDF -------------------------

    def _convert_via_libreoffice(self, input_path: str, output_dir: str) -> str:
        soffice = _soffice_yolu()
        if not soffice:
            raise DonusturmeHatasi(
                "LibreOffice bulunamadi. 'soffice' PATH'te degil ve su konumlarda yok: "
                + ", ".join(SOFFICE_ADAYLAR) +
                ". LibreOffice'i kurup PATH'e ekleyin ya da yukaridaki konumlardan birine kurun."
            )

        input_path = os.path.realpath(input_path)
        output_dir = os.path.realpath(output_dir)
        kaynak = Path(input_path)
        metin = kaynak.suffix.lower() in (".md", ".txt")
        if metin:
            try:
                kaynak.read_text(encoding="utf-8-sig")
            except UnicodeDecodeError as exc:
                raise DonusturmeHatasi("MD/TXT dosyasini UTF-8 kodlamasiyla kaydedin.") from exc

        hedef = os.path.realpath(os.path.join(output_dir, kaynak.stem + ".pdf"))
        # Ayri cikti: onceden kalan PDF basari sayilmaz. Ayri profil: acik
        # LibreOffice oturumuyla cakisma olmaz. Hata/zaman asiminda da temizlenir.
        # Profil sistemin gecici alaninda tutulur; cikti ise atomik tasima
        # icin hedefle ayni dosya sisteminde olmalidir.
        with tempfile.TemporaryDirectory(prefix="donusturucu_profil_") as profil_dir, \
                tempfile.TemporaryDirectory(prefix="donusturucu_", dir=output_dir) as gecici:
            calisma = Path(os.path.realpath(gecici))
            profil = Path(os.path.realpath(profil_dir)).as_uri()
            komut = [soffice, f"-env:UserInstallation={profil}", "--headless", "--norestore"]
            if metin:
                # Uzantiyi tahmin ettirmeyiz; eski LibreOffice surumlerinde de
                # MD duz metin acilir, HTML/Markdown olarak yorumlanmaz.
                komut.append("--infilter=Text (encoded):UTF8")
            komut.extend(["--convert-to", "pdf", "--outdir", str(calisma), input_path])
            try:
                sonuc = subprocess.run(
                    komut, shell=False, capture_output=True, text=True,
                    errors="replace", timeout=LIBREOFFICE_TIMEOUT, check=False,
                )
            except subprocess.TimeoutExpired as exc:
                raise DonusturmeHatasi(
                    f"LibreOffice {LIBREOFFICE_TIMEOUT} saniye icinde bitirmedi (zaman asimi)."
                ) from exc

            beklenen = Path(os.path.realpath(calisma / (kaynak.stem + ".pdf")))
            if sonuc.returncode != 0 or not beklenen.is_file() or beklenen.stat().st_size == 0:
                detay = (sonuc.stderr or "").strip() or (sonuc.stdout or "").strip() or "(ek bilgi yok)"
                raise DonusturmeHatasi(
                    f"LibreOffice donusumu basarisiz (cikis kodu: {sonuc.returncode}).\n"
                    f"LibreOffice ciktisi: {detay}"
                )
            # Basarili ciktiyi ayni dosya sisteminde tek adimda yerlestir.
            os.replace(beklenen, hedef)
        return hedef

    # --- PyMuPDF + python-pptx: PDF -> PPTX --------------------------------

    def _convert_pdf_to_pptx(self, input_path: str, output_path: str) -> None:
        if not fitz or not Presentation:
            raise DonusturmeHatasi("PyMuPDF ve python-pptx kütüphaneleri eksik.")

        prs = Presentation()
        doc = fitz.open(input_path)
        blank_layout = prs.slide_layouts[6]

        try:
            for page in doc:
                rect = page.rect
                prs.slide_width = Inches(rect.width / 72.0)
                prs.slide_height = Inches(rect.height / 72.0)
                slide = prs.slides.add_slide(blank_layout)
                pix = page.get_pixmap(dpi=150)
                slide.shapes.add_picture(
                    io.BytesIO(pix.tobytes("png")), 0, 0,
                    width=prs.slide_width, height=prs.slide_height,
                )
            prs.save(output_path)
        finally:
            doc.close()

    # --- PyMuPDF: PDF -> PNG ------------------------------------------------

    def _convert_pdf_to_png(self, input_path: str, output_dir: str) -> str:
        if not fitz:
            raise DonusturmeHatasi("PyMuPDF kütüphanesi eksik.")

        doc = fitz.open(input_path)
        base_name = os.path.splitext(os.path.basename(input_path))[0]
        first_output = None
        try:
            for idx, page in enumerate(doc, 1):
                pix = page.get_pixmap(dpi=150)
                filename = f"{base_name}.png" if len(doc) == 1 else f"{base_name}_sayfa_{idx}.png"
                out_path = os.path.join(output_dir, filename)
                pix.save(out_path)
                if idx == 1:
                    first_output = out_path
        finally:
            doc.close()
        return first_output

    # --- Pillow: PNG/JPG -> PDF ----------------------------------------------

    def _convert_image_to_pdf(self, input_path: str, output_path: str) -> None:
        with Image.open(input_path) as img:
            if img.mode != "RGB":
                img = img.convert("RGB")
            img.save(output_path, "PDF")

    # --- Genel giris noktasi --------------------------------------------------

    def convert_file(self, file_path: str, target_ext: str = None, output_dir: str = None) -> dict:
        abs_input = os.path.realpath(file_path)
        if not os.path.isfile(abs_input):
            # Mesaj modelin DUZELTEBILECEGI bilgiyi tasimali: hangi yol,
            # ve ayni klasorde benzer adlar. Bunun yoklugu bir kosuda
            # dort adim bosa harcadi ve model baska dosyayi ikame etti.
            benzer: list[str] = []
            dizin = os.path.dirname(abs_input)
            try:
                benzer = difflib.get_close_matches(
                    os.path.basename(abs_input), os.listdir(dizin), n=5, cutoff=0.5)
            except OSError:
                pass
            ipucu = f" Ayni klasordeki benzer adlar: {', '.join(benzer)}." if benzer else ""
            return {"success": False, "source": abs_input, "target": None,
                    "error": f"Dosya bulunamadi: {abs_input}.{ipucu}"}

        dir_name, full_name = os.path.split(abs_input)
        base_name, src_ext = os.path.splitext(full_name)
        src_ext = src_ext.lower()

        allowed = self.CONVERSION_MAP.get(src_ext, [])
        if not allowed:
            return {
                "success": False, "source": abs_input, "target": None,
                "error": (f"'{src_ext}' kaynak formati desteklenmiyor. Desteklenen kaynak "
                          f"formatlar: {', '.join(sorted(self.CONVERSION_MAP))}"),
            }

        target_ext = target_ext.lower() if target_ext else allowed[0]
        if not target_ext.startswith("."):
            target_ext = f".{target_ext}"

        if target_ext not in allowed:
            return {
                "success": False, "source": abs_input, "target": None,
                "error": (f"'{src_ext}' -> '{target_ext}' desteklenmiyor. '{src_ext}' icin "
                          f"desteklenen hedefler: {', '.join(allowed)}"),
            }

        target_dir = os.path.realpath(output_dir) if output_dir else dir_name
        target_path = os.path.realpath(os.path.join(target_dir, f"{base_name}{target_ext}"))

        try:
            os.makedirs(target_dir, exist_ok=True)
            if src_ext in (".pptx", ".ppt", ".docx", ".md", ".txt"):
                target_path = self._convert_via_libreoffice(abs_input, target_dir)
            elif src_ext == ".pdf" and target_ext == ".pptx":
                self._convert_pdf_to_pptx(abs_input, target_path)
            elif src_ext == ".pdf" and target_ext == ".png":
                target_path = self._convert_pdf_to_png(abs_input, target_dir)
            elif src_ext in (".png", ".jpg", ".jpeg") and target_ext == ".pdf":
                self._convert_image_to_pdf(abs_input, target_path)
            else:
                return {
                    "success": False, "source": abs_input, "target": None,
                    "error": f"'{src_ext}' -> '{target_ext}' icin bir donusturucu tanimli degil.",
                }
            return {"success": True, "source": abs_input, "target": target_path, "error": None}
        except DonusturmeHatasi as e:
            return {"success": False, "source": abs_input, "target": None, "error": str(e)}
        except Exception as e:
            return {"success": False, "source": abs_input, "target": None, "error": f"{type(e).__name__}: {e}"}
