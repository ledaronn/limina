# belge.py — PDF/DOCX/PPTX/XLSX dosyalarını düz metne çevirir. Sadece okur.
#
# Goruntuler ve metin katmani olmayan (taranmis) PDF sayfalari OCR'dan gecer
# (pevrai/ocr.py, Windows yerlesik OCR). OCR kullanilamiyorsa neden metinde
# soylenir; arac yine calisir, istisna sizmaz.
from __future__ import annotations
from pevrai.sonuc import hata as sonuc_hatasi

from pathlib import Path

from pevrai import ocr
from pevrai.ceviri import t

DESTEKLENEN = {".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md", ".csv"} | ocr.GORUNTU_UZANTILARI


def oku(yol: Path, sayfa_bas: int = 1, sayfa_son: int = 0) -> str:
    """Dosya türüne göre uygun okuyucuya yönlendirir. Her zaman metin döndürür."""
    uzanti = yol.suffix.lower()
    try:
        if uzanti == ".pdf":
            return _pdf(yol, sayfa_bas, sayfa_son)
        if uzanti == ".docx":
            return _docx(yol)
        if uzanti == ".pptx":
            return _pptx(yol)
        if uzanti == ".xlsx":
            return _xlsx(yol)
        if uzanti in {".txt", ".md", ".csv"}:
            return yol.read_text(encoding="utf-8", errors="replace")
        if uzanti in ocr.GORUNTU_UZANTILARI:
            return ocr.goruntu_oku(yol)
    except ImportError as e:
        return sonuc_hatasi(t("Bu dosya türü için gerekli kütüphane kurulu değil: {hata}", hata=e))
    except Exception as e:
        return sonuc_hatasi(t("'{ad}' okunamadı ({tur}): {hata}", ad=yol.name, tur=type(e).__name__, hata=e))

    return sonuc_hatasi(t("'{uzanti}' uzantısı okunamıyor. Desteklenenler: {liste}. "
             "Bu bir sunum/belge değilse önce dönüştür.", uzanti=uzanti, liste=", ".join(sorted(DESTEKLENEN))))


def _pdf(yol: Path, bas: int, son: int) -> str:
    import pymupdf as fitz

    with fitz.open(yol) as doc:
        toplam = len(doc)
        bitis = son if son and son <= toplam else toplam
        baslangic = max(1, bas)
        if baslangic > toplam:
            return sonuc_hatasi(t("'{ad}' {toplam} sayfa. İstenen başlangıç sayfası ({bas}) dışarıda.",
                     ad=yol.name, toplam=toplam, bas=baslangic))

        parcalar = [t("[PDF: {ad}, {toplam} sayfa, {bas}-{son} arası gösteriliyor]",
                      ad=yol.name, toplam=toplam, bas=baslangic, son=bitis)]
        # Metin katmani olmayan sayfa taranmis demektir: OCR'a duser. Sayfa
        # basina karar verilir — karisik belgelerde (kapak taranmis, icerik
        # metin) yalnizca gereken sayfa OCR'lanir. OCR'lanan sayfa basliginda
        # isaretlenir: model "bu metin OCR'dan, hatali karakter olabilir" bilir.
        ocr_yok = ocr.kullanilamaz_nedeni()
        bos_sayfa = ocr_sayfa = 0
        for i in range(baslangic - 1, bitis):
            metin = doc[i].get_text().strip()
            etiket = ""
            if not metin and not ocr_yok and ocr_sayfa < ocr.AZAMI_SAYFA:
                metin = ocr.pdf_sayfasi_oku(doc[i])
                ocr_sayfa += 1
                etiket = " [OCR]"
            if not metin:
                bos_sayfa += 1
            parcalar.append(f"\n--- sayfa {i + 1}{etiket} ---\n{metin}")

        if bos_sayfa == bitis - baslangic + 1:
            if ocr_yok:
                return sonuc_hatasi(t("'{ad}' içinde metin katmanı yok — muhtemelen taranmış bir belge (sayfalar "
                         "görüntü olarak saklanmış). Metni çıkarmak için OCR gerekir. {neden}",
                         ad=yol.name, neden=ocr_yok))
            return sonuc_hatasi(t("'{ad}' içinde metin katmanı yok ve OCR da bu sayfalarda metin bulamadı "
                     "(boş sayfalar ya da çok düşük kaliteli tarama olabilir).", ad=yol.name))
        if ocr_sayfa:
            parcalar.insert(1, t("[{n} sayfa OCR ile okundu: hatalı karakter olabilir, sayı ve adları "
                                 "doğrulamadan kullanma]", n=ocr_sayfa))
        return "\n".join(parcalar)


def _docx(yol: Path) -> str:
    import docx

    d = docx.Document(str(yol))
    satirlar = [p.text for p in d.paragraphs if p.text.strip()]
    for tablo in d.tables:
        for satir in tablo.rows:
            hucreler = [h.text.strip() for h in satir.cells]
            if any(hucreler):
                satirlar.append(" | ".join(hucreler))
    return f"[DOCX: {yol.name}]\n" + ("\n".join(satirlar) or t("(belge boş)"))


def _pptx(yol: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(yol))
    parcalar = [f"[PPTX: {yol.name}, {len(prs.slides)} slayt]"]
    for i, slayt in enumerate(prs.slides, 1):
        metinler = [s.text.strip() for s in slayt.shapes if s.has_text_frame and s.text.strip()]
        parcalar.append(f"\n--- slayt {i} ---\n" + ("\n".join(metinler) or "(metin yok)"))
    return "\n".join(parcalar)


def _xlsx(yol: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(yol), data_only=True, read_only=True)
    parcalar = [f"[XLSX: {yol.name}, sayfalar: {', '.join(wb.sheetnames)}]"]
    for ad in wb.sheetnames:
        ws = wb[ad]
        parcalar.append(f"\n--- {ad} ---")
        for satir in ws.iter_rows(max_row=200, values_only=True):
            hucreler = [str(h) for h in satir if h is not None]
            if hucreler:
                parcalar.append(" | ".join(hucreler))
    wb.close()
    return "\n".join(parcalar)
